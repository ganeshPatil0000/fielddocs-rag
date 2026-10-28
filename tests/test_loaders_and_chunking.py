import io

import pytest

from app.ingestion.chunking import chunk_documents, chunk_ids, clean_text
from app.ingestion.loaders import (
    EmptyDocumentError,
    UnsupportedFileError,
    load_document,
)


def _minimal_pdf(text: str) -> bytes:
    """Build a valid one-page PDF with a text layer (no external libraries)."""
    stream = f"BT /F1 12 Tf 72 712 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def test_pdf_loader_keeps_page_numbers():
    docs = load_document(_minimal_pdf("Received power must exceed -27 dBm"), "guide.pdf")
    assert len(docs) == 1
    assert docs[0].metadata == {"source": "guide.pdf", "page": 1}
    assert "-27 dBm" in docs[0].page_content


def test_docx_loader_marks_headings_and_tables():
    import docx

    d = docx.Document()
    d.add_heading("Battery alarm", level=2)
    d.add_paragraph("Replace batteries every 4 years.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Float voltage"
    table.rows[0].cells[1].text = "52-54 V"
    buf = io.BytesIO()
    d.save(buf)

    docs = load_document(buf.getvalue(), "runbook.docx")
    text = docs[0].page_content
    assert "## Battery alarm" in text
    assert "Float voltage | 52-54 V" in text


def test_unsupported_and_empty_files_are_rejected():
    with pytest.raises(UnsupportedFileError):
        load_document(b"data", "virus.exe")
    with pytest.raises(EmptyDocumentError):
        load_document(b"   \n ", "empty.txt")


def test_clean_text_rejoins_hyphenation_and_collapses_space():
    assert clean_text("instal-\nlation   guide\n\n\n\nnext") == "installation guide\n\nnext"


def test_chunks_carry_heading_path_and_stable_metadata():
    md = "# Guide\n\n## LOS fault\n\n" + "Check the patch lead is seated. " * 60
    docs = load_document(md.encode(), "fttp.md")
    chunks = chunk_documents(docs, "doc1", chunk_size=400, chunk_overlap=50)

    assert len(chunks) > 1
    assert all(c.page_content.startswith("[fttp.md > Guide > LOS fault]") for c in chunks)
    assert [c.metadata["chunk_index"] for c in chunks] == list(range(len(chunks)))
    assert all(c.metadata["doc_id"] == "doc1" for c in chunks)
    # Same input -> same hashes (idempotent re-ingestion).
    again = chunk_documents(docs, "doc1", chunk_size=400, chunk_overlap=50)
    assert [c.metadata["chunk_hash"] for c in chunks] == [c.metadata["chunk_hash"] for c in again]


def test_chunk_ids_are_deterministic():
    assert chunk_ids("abc", 3) == ["abc#0", "abc#1", "abc#2"]
    assert chunk_ids("abc", 5, start=3) == ["abc#3", "abc#4"]
