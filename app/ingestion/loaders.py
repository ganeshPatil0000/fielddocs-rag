"""Turn raw file bytes into LangChain Documents (one per page / section)."""

from __future__ import annotations

import io
from pathlib import PurePosixPath

from langchain_core.documents import Document


class IngestionError(Exception):
    """Base class. ``retryable`` tells the queue worker whether to retry."""

    retryable = False


class UnsupportedFileError(IngestionError):
    pass


class EmptyDocumentError(IngestionError):
    pass


def extension(filename: str) -> str:
    return PurePosixPath(filename).suffix.lower()


def _load_pdf(data: bytes, filename: str) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:  # noqa: BLE001
            raise UnsupportedFileError("PDF is password protected") from exc
    docs = []
    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(Document(page_content=text, metadata={"source": filename, "page": page_no}))
    return docs


def _load_docx(data: bytes, filename: str) -> list[Document]:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        # Keep headings visible as markdown so the chunker can use them.
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            parts.append(f"{'#' * min(int(level), 3)} {text}")
        else:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = "\n\n".join(parts)
    return [Document(page_content=text, metadata={"source": filename})] if text.strip() else []


def _load_text(data: bytes, filename: str) -> list[Document]:
    text = data.decode("utf-8", errors="replace")
    return [Document(page_content=text, metadata={"source": filename})] if text.strip() else []


LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".txt": _load_text,
    ".md": _load_text,
}


def load_document(data: bytes, filename: str) -> list[Document]:
    ext = extension(filename)
    loader = LOADERS.get(ext)
    if loader is None:
        raise UnsupportedFileError(f"Unsupported file type '{ext}'")
    docs = loader(data, filename)
    if not docs:
        hint = " (scanned PDF? add OCR, e.g. Amazon Textract)" if ext == ".pdf" else ""
        raise EmptyDocumentError(f"No extractable text in {filename}{hint}")
    return docs
